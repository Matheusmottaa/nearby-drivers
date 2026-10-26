"""Gera a apresentação do nearby-drivers.

O arquivo de saída é um .pptx. No Google Drive, abra com Google Slides.
"""

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt
from lxml import etree

# Paleta: papel quente, tinta e âmbar. Georgia + Arial existem no Google Slides.
INK = RGBColor(0x0F, 0x1C, 0x2E)
INK_SOFT = RGBColor(0x1A, 0x32, 0x4A)
PAPER = RGBColor(0xF6, 0xF3, 0xEC)
WHITE = RGBColor(0xFF, 0xFC, 0xF8)
AMBER = RGBColor(0xE0, 0x9A, 0x2B)
AMBER_DEEP = RGBColor(0xA8, 0x6E, 0x12)
TEAL = RGBColor(0x1C, 0x74, 0x66)
CORAL = RGBColor(0xC4, 0x5C, 0x4A)
SLATE = RGBColor(0x5C, 0x6A, 0x7A)
BODY = RGBColor(0x1C, 0x28, 0x36)
LINE = RGBColor(0xE4, 0xDD, 0xD2)
MUTED = RGBColor(0xC5, 0xD0, 0xDB)
AMBER_WASH = RGBColor(0xFB, 0xF3, 0xE0)
TEAL_WASH = RGBColor(0xE6, 0xF2, 0xEF)
CORAL_WASH = RGBColor(0xF8, 0xEC, 0xE8)
INK_WASH = RGBColor(0xEE, 0xF2, 0xF6)

SERIF = "Georgia"
SANS = "Arial"

W = Inches(13.333333)
H = Inches(7.5)
TOTAL = 16


def rgb_hex(color):
    return f"{color[0]:02X}{color[1]:02X}{color[2]:02X}"


def set_alpha(shape, opacity):
    """opacity de 0 a 1 sobre o preenchimento sólido já definido."""
    solid = shape._element.spPr.find(qn("a:solidFill"))
    if solid is None:
        return
    srgb = solid.find(qn("a:srgbClr"))
    if srgb is None:
        return
    for child in list(srgb):
        if child.tag == qn("a:alpha"):
            srgb.remove(child)
    alpha = etree.SubElement(srgb, qn("a:alpha"))
    alpha.set("val", str(int(opacity * 100000)))


def no_line(shape):
    shape.line.fill.background()


def fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    no_line(shape)


def rect(slide, x, y, w, h, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    fill(shape, color)
    return shape


def round_rect(slide, x, y, w, h, color, radius=0.12, line_color=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    shape.adjustments[0] = radius
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    if line_color is None:
        no_line(shape)
    else:
        shape.line.color.rgb = line_color
        shape.line.width = Pt(1)
    return shape


def oval(slide, x, y, d, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, x, y, d, d)
    fill(shape, color)
    return shape


def style_run(run, font, size, color, bold=False, italic=False, tracking=None):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    if tracking is not None:
        run._r.get_or_add_rPr().set("spc", str(tracking))


def add_text(
    slide,
    x,
    y,
    w,
    h,
    text,
    font,
    size,
    color,
    bold=False,
    italic=False,
    align=PP_ALIGN.LEFT,
    anchor=MSO_ANCHOR.TOP,
    tracking=None,
):
    box = slide.shapes.add_textbox(x, y, w, h)
    frame = box.text_frame
    frame.word_wrap = True
    frame.auto_size = None
    frame.anchor = anchor
    frame.margin_left = Emu(0)
    frame.margin_right = Emu(0)
    frame.margin_top = Emu(0)
    frame.margin_bottom = Emu(0)
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.space_before = Pt(0)
    paragraph.space_after = Pt(0)
    run = paragraph.add_run()
    run.text = text
    style_run(run, font, size, color, bold, italic, tracking)
    return box


def add_lines(slide, x, y, w, h, lines, font, size, color, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.0, after=0):
    box = slide.shapes.add_textbox(x, y, w, h)
    frame = box.text_frame
    frame.word_wrap = True
    frame.auto_size = None
    frame.anchor = anchor
    frame.margin_left = Emu(0)
    frame.margin_right = Emu(0)
    frame.margin_top = Emu(0)
    frame.margin_bottom = Emu(0)
    for i, line in enumerate(lines):
        paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        paragraph.alignment = align
        paragraph.space_before = Pt(0)
        paragraph.space_after = Pt(after)
        paragraph.line_spacing = spacing
        run = paragraph.add_run()
        run.text = line
        style_run(run, font, size, color, bold)
    return box


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def light_chrome(slide, number):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = PAPER
    rect(slide, 0, 0, Inches(0.08), H, AMBER)
    rect(slide, Inches(0.62), Inches(7.08), Inches(12.1), Inches(0.01), LINE)
    add_text(slide, Inches(0.62), Inches(7.16), Inches(4), Inches(0.26), "nearby-drivers", SANS, 11, SLATE)
    add_text(
        slide,
        Inches(11.3),
        Inches(7.16),
        Inches(1.45),
        Inches(0.26),
        f"{number:02d}",
        SANS,
        11,
        SLATE,
        align=PP_ALIGN.RIGHT,
    )


def dark_chrome(slide, number):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = INK
    add_text(slide, Inches(0.72), Inches(7.05), Inches(4), Inches(0.26), "nearby-drivers", SANS, 12, MUTED)
    add_text(
        slide,
        Inches(11.2),
        Inches(7.05),
        Inches(1.4),
        Inches(0.26),
        f"{number:02d}",
        SANS,
        12,
        MUTED,
        align=PP_ALIGN.RIGHT,
    )


def kicker(slide, text, y=Inches(0.38)):
    rect(slide, Inches(0.62), y + Inches(0.06), Inches(0.28), Inches(0.035), AMBER)
    add_text(slide, Inches(1.02), y, Inches(8), Inches(0.28), text, SANS, 12, AMBER_DEEP, bold=True, tracking=180)


def title(slide, text, y=Inches(0.72), size=34, width=Inches(12)):
    add_text(slide, Inches(0.62), y, width, Inches(0.95), text, SERIF, size, INK)


def card(slide, x, y, w, h, wash=WHITE):
    return round_rect(slide, x, y, w, h, wash, radius=0.06, line_color=LINE)


def accent(slide, x, y, h, color):
    """Barra curta dentro do cartão, sem brigar com o canto arredondado."""
    return rect(slide, x, y, Inches(0.07), h, color)


def slide_cover(prs):
    slide = blank(prs)
    dark_chrome(slide, 1)
    rect(slide, 0, 0, Inches(0.08), H, AMBER)

    add_text(slide, Inches(0.78), Inches(1.55), Inches(6), Inches(0.3), "SISTEMA DE CORRIDA", SANS, 13, AMBER, bold=True, tracking=220)
    add_text(slide, Inches(0.74), Inches(1.95), Inches(8), Inches(1.05), "nearby-drivers", SERIF, 60, WHITE)
    add_lines(
        slide,
        Inches(0.78),
        Inches(3.2),
        Inches(7.2),
        Inches(1.4),
        ["O motorista disponível", "mais próximo do embarque."],
        SERIF,
        26,
        MUTED,
        spacing=1.05,
    )
    rect(slide, Inches(0.78), Inches(4.85), Inches(1.35), Inches(0.045), AMBER)

    labels = ["ride", "driver", "matching", "routing"]
    x = 0.78
    for label in labels:
        oval(slide, Inches(x), Inches(5.25), Inches(0.14), AMBER)
        add_text(slide, Inches(x + 0.26), Inches(5.18), Inches(1.7), Inches(0.3), label, SANS, 15, WHITE)
        x += 1.85

    # Mapa estilizado: uma busca a partir do embarque.
    ox, oy = 8.55, 1.7
    ring = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(ox + 0.86), Inches(oy + 0.56), Inches(2.7), Inches(2.7))
    ring.fill.background()
    ring.line.color.rgb = AMBER
    ring.line.width = Pt(1.25)

    nodes = {
        "a": (0.35, 0.45),
        "b": (0.55, 2.15),
        "c": (0.25, 3.55),
        "d": (1.85, 1.15),
        "pin": (2.15, 1.85),
        "e": (2.05, 3.35),
        "f": (3.55, 0.7),
        "g": (3.7, 2.45),
        "h": (3.35, 3.9),
    }
    edges = [("a", "d"), ("b", "d"), ("b", "pin"), ("c", "e"), ("d", "pin"), ("pin", "e"), ("pin", "f"), ("pin", "g"), ("e", "h"), ("g", "h"), ("f", "g")]
    for u, v in edges:
        x1, y1 = nodes[u]
        x2, y2 = nodes[v]
        line = slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT,
            Inches(ox + x1 + 0.11),
            Inches(oy + y1 + 0.11),
            Inches(ox + x2 + 0.11),
            Inches(oy + y2 + 0.11),
        )
        line.line.color.rgb = RGBColor(0x3E, 0x5C, 0x78)
        line.line.width = Pt(1.5)

    for name, (x, y) in nodes.items():
        if name == "pin":
            continue
        dot = oval(slide, Inches(ox + x), Inches(oy + y), Inches(0.22), RGBColor(0xD5, 0xE2, 0xEE))
        set_alpha(dot, 0.95)
    halo = oval(slide, Inches(ox + 1.95), Inches(oy + 1.65), Inches(0.52), AMBER)
    set_alpha(halo, 0.35)
    oval(slide, Inches(ox + 2.08), Inches(oy + 1.78), Inches(0.26), AMBER)

    notes(
        slide,
        "Abre com a promessa do sistema: dado um embarque, achar o motorista livre que chega antes e dizer em quanto tempo. "
        "Os quatro nomes embaixo são as aplicações. O desenho é o embarque no centro e motoristas ao redor, ligados pelas ruas.",
    )


def slide_promise(prs):
    slide = blank(prs)
    light_chrome(slide, 2)
    kicker(slide, "O QUE O PASSAGEIRO PEDE")
    add_lines(
        slide,
        Inches(0.62),
        Inches(1.85),
        Inches(12),
        Inches(3.3),
        [
            "Um ponto de embarque.",
            "O motorista livre que chega antes.",
            "Quem vem, e em quanto tempo.",
        ],
        SERIF,
        36,
        INK,
        spacing=1.2,
    )
    rect(slide, Inches(0.64), Inches(5.15), Inches(1.15), Inches(0.04), AMBER)
    add_text(
        slide,
        Inches(0.64),
        Inches(5.4),
        Inches(10),
        Inches(0.8),
        "A corrida termina quando isso está respondido.",
        SANS,
        18,
        SLATE,
    )
    notes(
        slide,
        "A corrida termina em motorista atribuído. Não há acompanhamento da viagem neste sistema: "
        "o sucesso é o passageiro saber quem está a caminho e qual ETA foi prometido.",
    )


def slide_apps(prs):
    slide = blank(prs)
    light_chrome(slide, 3)
    kicker(slide, "AS APLICAÇÕES")
    title(slide, "Quatro trabalhos, quatro donos")
    add_text(
        slide,
        Inches(0.64),
        Inches(1.62),
        Inches(11),
        Inches(0.4),
        "Nenhuma aplicação faz o trabalho da outra.",
        SANS,
        16,
        SLATE,
    )

    apps = [
        ("RIDE-SERVICE", "A corrida", "Cria o pedido, guarda o estado e confirma o motorista.", "1 instância", AMBER, AMBER_WASH),
        ("DRIVER-SERVICE", "A posição", "Sabe quais motoristas estão perto e disponíveis.", "1 instância", TEAL, TEAL_WASH),
        ("MATCHING-SERVICE", "A lista", "Ordena os candidatos por quem chega antes.", "de 1 a 3", INK_SOFT, INK_WASH),
        ("ROUTING-SERVICE", "O tempo", "Calcula quanto cada um leva pelas ruas.", "de 2 a 4", CORAL, CORAL_WASH),
    ]
    positions = [(0.62, 2.25), (6.95, 2.25), (0.62, 4.55), (6.95, 4.55)]
    for (label, name, body, meta, tone, wash), (x, y) in zip(apps, positions):
        card(slide, Inches(x), Inches(y), Inches(5.75), Inches(2.1), WHITE)
        accent(slide, Inches(x + 0.22), Inches(y + 0.28), Inches(1.54), tone)
        add_text(slide, Inches(x + 0.46), Inches(y + 0.22), Inches(4.9), Inches(0.26), label, SANS, 12, tone, bold=True, tracking=120)
        add_text(slide, Inches(x + 0.46), Inches(y + 0.52), Inches(5.0), Inches(0.42), name, SERIF, 24, INK)
        add_text(slide, Inches(x + 0.46), Inches(y + 1.08), Inches(5.0), Inches(0.55), body, SANS, 15, BODY)
        add_text(slide, Inches(x + 0.46), Inches(y + 1.66), Inches(3), Inches(0.28), meta, SANS, 12, SLATE)

    notes(
        slide,
        "Ride é a dona do estado. Driver responde quem está perto. Matching só monta o ranking. "
        "Routing só devolve tempo e distância. Ride e driver ficam em uma instância: cada um guarda estado. "
        "Matching sobe de 1 a 3 e routing de 2 a 4 quando a CPU passa de 60%.",
    )


def slide_flow(prs):
    slide = blank(prs)
    light_chrome(slide, 4)
    kicker(slide, "O CAMINHO DE UMA CORRIDA")
    title(slide, "O aceite não espera a busca", size=32)
    add_text(
        slide,
        Inches(0.64),
        Inches(1.62),
        Inches(11.5),
        Inches(0.4),
        "O passageiro recebe a confirmação na hora. O resto acontece por mensagens.",
        SANS,
        16,
        SLATE,
    )

    steps = [
        ("1", "Aceita", "A corrida entra no banco e a resposta sai na hora."),
        ("2", "Procura", "Motoristas livres num raio de 3 km."),
        ("3", "Mede", "O tempo de cada um até o embarque."),
        ("4", "Reserva", "O primeiro que ainda está livre."),
        ("5", "Entrega", "Quem vem, e em quanto tempo."),
    ]
    x = 0.62
    width = 2.28
    gap = 0.16
    for number, name, body in steps:
        card(slide, Inches(x), Inches(2.35), Inches(width), Inches(4.15), WHITE)
        add_text(slide, Inches(x + 0.2), Inches(2.55), Inches(1.8), Inches(0.6), number, SERIF, 32, AMBER)
        add_text(slide, Inches(x + 0.2), Inches(3.35), Inches(1.9), Inches(0.45), name, SERIF, 22, INK)
        add_text(slide, Inches(x + 0.2), Inches(4.0), Inches(1.9), Inches(1.9), body, SANS, 14, BODY)
        x += width + gap

    notes(
        slide,
        "POST /rides responde 202 em REQUESTED. O matching consome rides.requested, chama nearby e depois eta em lote, "
        "e publica o ranking. O ride reserva o primeiro livre e publica rides.assigned. O driver marca o motorista como em corrida.",
    )


def slide_ride(prs):
    slide = blank(prs)
    light_chrome(slide, 5)
    kicker(slide, "APLICAÇÃO  01")
    title(slide, "ride-service")
    add_text(
        slide,
        Inches(0.64),
        Inches(1.62),
        Inches(11),
        Inches(0.4),
        "A única que pode criar, atribuir ou cancelar uma corrida.",
        SANS,
        18,
        SLATE,
    )

    points = [
        ("Responde na hora", "Confirma o pedido antes de procurar o motorista. O passageiro não fica preso na busca."),
        ("Estado no banco", "A corrida e o aviso saem na mesma transação. Se o processo cair no meio, o aviso é reenviado."),
        ("Ela decide de verdade", "O matching sugere uma lista. Aqui se vê se o motorista ainda está livre e se grava o resultado."),
    ]
    y = 2.25
    for name, body in points:
        card(slide, Inches(0.62), Inches(y), Inches(12.1), Inches(1.4), WHITE)
        accent(slide, Inches(0.9), Inches(y + 0.32), Inches(0.76), AMBER)
        add_text(slide, Inches(1.2), Inches(y + 0.22), Inches(11.1), Inches(0.38), name, SERIF, 20, INK)
        add_text(slide, Inches(1.2), Inches(y + 0.7), Inches(11.1), Inches(0.5), body, SANS, 15, BODY)
        y += 1.55

    notes(
        slide,
        "Não há chamada REST desta aplicação para as outras. A outbox garante que o evento rides.requested existe junto com a corrida. "
        "A reserva rápida no Redis falha cedo; a garantia final é o índice único de motorista atribuído no Postgres.",
    )


def slide_states(prs):
    slide = blank(prs)
    light_chrome(slide, 6)
    kicker(slide, "RIDE-SERVICE")
    title(slide, "A vida de uma corrida", size=32)

    states = [
        ("Pedido", "A corrida já existe no banco.", AMBER),
        ("Buscando", "A lista de motoristas está sendo montada.", TEAL),
        ("A caminho", "O passageiro sabe quem vem e o tempo.", INK_SOFT),
    ]
    card_w = 3.35
    arrow_w = 0.42
    gap = 0.18
    total = 3 * card_w + 2 * arrow_w + 4 * gap
    x = (13.333 - total) / 2
    for i, (name, body, accent) in enumerate(states):
        card(slide, Inches(x), Inches(1.9), Inches(card_w), Inches(2.55), WHITE)
        rect(slide, Inches(x + 0.28), Inches(2.18), Inches(0.42), Inches(0.055), accent)
        add_text(slide, Inches(x + 0.28), Inches(2.38), Inches(card_w - 0.5), Inches(0.55), name, SERIF, 26, INK)
        add_text(slide, Inches(x + 0.28), Inches(3.0), Inches(card_w - 0.5), Inches(1.1), body, SANS, 15, BODY)
        x += card_w + gap
        if i < 2:
            add_text(slide, Inches(x - 0.02), Inches(2.75), Inches(arrow_w), Inches(0.5), "→", SERIF, 26, AMBER, align=PP_ALIGN.CENTER)
            x += arrow_w + gap

    outcomes = [
        (CORAL, "Cancelada", "O passageiro pode sair enquanto espera."),
        (SLATE, "Sem motorista", "Ninguém no raio, ou a lista já foi toda reservada."),
        (AMBER, "Motorista sumiu", "A busca recomeça sem ele. No máximo, 3 vezes."),
    ]
    x = 0.62
    for color, name, body in outcomes:
        card(slide, Inches(x), Inches(4.75), Inches(3.9), Inches(1.9), WHITE)
        oval(slide, Inches(x + 0.28), Inches(5.0), Inches(0.16), color)
        add_text(slide, Inches(x + 0.58), Inches(4.92), Inches(3.05), Inches(0.35), name, SANS, 15, INK, bold=True)
        add_text(slide, Inches(x + 0.28), Inches(5.45), Inches(3.35), Inches(0.95), body, SANS, 14, BODY)
        x += 4.08

    notes(
        slide,
        "Estados: REQUESTED, MATCHING, DRIVER_ASSIGNED. Terminais de falha: CANCELLED e NO_DRIVER_AVAILABLE. "
        "Se o motorista reservado para de mandar heartbeat, a corrida volta a MATCHING, publica release e um novo requested sem aquele motorista. "
        "RIDE_MAX_ATTEMPTS padrão é 3. Rematch conta como tentativa.",
    )


def slide_driver(prs):
    slide = blank(prs)
    light_chrome(slide, 7)
    kicker(slide, "APLICAÇÃO  02")
    title(slide, "driver-service")
    add_text(
        slide,
        Inches(0.64),
        Inches(1.58),
        Inches(11.5),
        Inches(0.38),
        "Cadastro no banco. Posição na memória. Não escolhe motorista.",
        SANS,
        16,
        SLATE,
    )

    columns = [
        (TEAL, "Cadastro", "Muda pouco: nome, placa, classe, status. Fica no PostgreSQL e não pode se perder."),
        (AMBER, "Posição", "Chega o tempo todo. Latitude, longitude e um número de sequência. O banco não guarda coordenada."),
    ]
    x = 0.62
    for color, name, body in columns:
        card(slide, Inches(x), Inches(2.15), Inches(6.0), Inches(2.15), WHITE)
        accent(slide, Inches(x + 0.28), Inches(2.48), Inches(1.5), color)
        add_text(slide, Inches(x + 0.56), Inches(2.38), Inches(5.15), Inches(0.45), name, SERIF, 24, INK)
        add_text(slide, Inches(x + 0.56), Inches(3.0), Inches(5.15), Inches(1.0), body, SANS, 15, BODY)
        x += 6.2

    card(slide, Inches(0.62), Inches(4.55), Inches(12.1), Inches(2.1), WHITE)
    add_text(slide, Inches(0.92), Inches(4.75), Inches(11.4), Inches(0.36), "Entra na busca só se as três valem ao mesmo tempo", SERIF, 18, INK)
    pills = [
        (TEAL_WASH, TEAL, "Disponível"),
        (AMBER_WASH, AMBER_DEEP, "Falou nos últimos 30 s"),
        (INK_WASH, INK_SOFT, "Não está em corrida"),
    ]
    x = 0.92
    for wash, color, label in pills:
        round_rect(slide, Inches(x), Inches(5.4), Inches(3.55), Inches(0.78), wash, radius=0.4)
        add_text(slide, Inches(x), Inches(5.58), Inches(3.55), Inches(0.42), label, SANS, 15, color, bold=True, align=PP_ALIGN.CENTER)
        x += 3.8

    notes(
        slide,
        "Disponível significa status AVAILABLE, heartbeat nos últimos 30 segundos e não reservado. "
        "O bit vive no slot em memória, então um app que morreu sai da busca sem UPDATE no Postgres. "
        "PUT /location com seq menor ou igual é ignorado. No boot, o replay dos últimos 5 minutos de drivers.location reconstitui as posições.",
    )


def slide_kdtree(prs):
    slide = blank(prs)
    light_chrome(slide, 8)
    kicker(slide, "ALGORITMO")
    title(slide, "Achar quem está perto", size=32)
    add_text(
        slide,
        Inches(0.64),
        Inches(1.62),
        Inches(6.6),
        Inches(0.7),
        "Uma KD-Tree organiza as posições para não medir a cidade inteira.",
        SANS,
        16,
        SLATE,
    )

    points = [
        ("Corta o mapa", "A árvore divide a cidade em regiões. A consulta descarta as que estão longe do embarque."),
        ("Devolve os mais perto", "Até 10 motoristas dentro de 3 km, em linha reta. A ordem final, por tempo, vem depois."),
        ("Não trava a leitura", "A árvore é refeita por trás. Quem consulta usa a cópia pronta, mais os GPS que chegaram depois."),
    ]
    y = 2.4
    for name, body in points:
        add_text(slide, Inches(0.64), Inches(y), Inches(6.7), Inches(0.34), name, SANS, 16, INK, bold=True)
        add_text(slide, Inches(0.64), Inches(y + 0.36), Inches(6.7), Inches(0.7), body, SANS, 14, BODY)
        y += 1.2

    # Diagrama: partições e a região do embarque.
    ox, oy, s = 8.05, 1.85, 4.35
    parts = [
        (0.00, 0.00, 0.46, 0.38, RGBColor(0xE4, 0xDB, 0xCC)),
        (0.00, 0.38, 0.46, 0.62, RGBColor(0xD5, 0xE3, 0xDF)),
        (0.46, 0.00, 0.28, 0.58, RGBColor(0xFD, 0xF6, 0xE4)),
        (0.74, 0.00, 0.26, 0.58, RGBColor(0xE4, 0xDB, 0xCC)),
        (0.46, 0.58, 0.54, 0.42, RGBColor(0xD5, 0xE0, 0xE8)),
    ]
    for px, py, pw, ph, color in parts:
        rect(slide, Inches(ox + px * s), Inches(oy + py * s), Inches(pw * s), Inches(ph * s), color)
    frame = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(ox), Inches(oy), Inches(s), Inches(s))
    frame.fill.background()
    frame.line.color.rgb = RGBColor(0xD9, 0xD0, 0xC2)
    frame.line.width = Pt(1.25)

    gray_dots = [(0.16, 0.18), (0.3, 0.72), (0.84, 0.22), (0.62, 0.78)]
    for px, py in gray_dots:
        oval(slide, Inches(ox + px * s), Inches(oy + py * s), Inches(0.16), RGBColor(0x9A, 0xA8, 0xB4))
    for px, py in ((0.55, 0.22), (0.66, 0.4)):
        oval(slide, Inches(ox + px * s), Inches(oy + py * s), Inches(0.18), TEAL)
    oval(slide, Inches(ox + 0.58 * s), Inches(oy + 0.3 * s), Inches(0.22), AMBER)

    add_text(
        slide,
        Inches(8.05),
        Inches(6.32),
        Inches(4.4),
        Inches(0.55),
        "A região clara é a única que a consulta percorre.",
        SANS,
        12,
        SLATE,
    )

    notes(
        slide,
        "A árvore é 2D sobre metros locais, para o corte não distorcer com a latitude. A distância devolvida ao chamador é Haversine. "
        "O writer atualiza o slot e anexa um delta. Um rebuild em background publica outro snapshot. "
        "Índice ainda não publicado responde 503, o mesmo sinal do readyz.",
    )


def slide_matching(prs):
    slide = blank(prs)
    light_chrome(slide, 9)
    kicker(slide, "APLICAÇÃO  03")
    title(slide, "matching-service", size=32)
    add_text(
        slide,
        Inches(0.64),
        Inches(1.58),
        Inches(7.2),
        Inches(0.55),
        "Não grava corrida, não reserva motorista e não calcula rota.",
        SANS,
        15,
        SLATE,
    )

    steps = [
        ("1", "Recebe o pedido", "Depois que a corrida já foi aceita."),
        ("2", "Busca candidatos", "Quem está livre, perto, na classe pedida."),
        ("3", "Pede o tempo", "Todos de uma vez, numa chamada só."),
        ("4", "Publica a lista", "Do mais rápido ao mais lento."),
    ]
    rect(slide, Inches(0.86), Inches(2.45), Inches(0.015), Inches(3.55), AMBER)
    y = 2.3
    for number, name, body in steps:
        oval(slide, Inches(0.72), Inches(y + 0.02), Inches(0.32), AMBER)
        add_text(
            slide,
            Inches(0.72),
            Inches(y + 0.02),
            Inches(0.32),
            Inches(0.32),
            number,
            SANS,
            12,
            WHITE,
            bold=True,
            align=PP_ALIGN.CENTER,
            anchor=MSO_ANCHOR.MIDDLE,
        )
        add_text(slide, Inches(1.25), Inches(y), Inches(5.6), Inches(0.32), name, SANS, 16, INK, bold=True)
        add_text(slide, Inches(1.25), Inches(y + 0.34), Inches(5.6), Inches(0.32), body, SANS, 14, SLATE)
        y += 0.95

    card(slide, Inches(8.15), Inches(2.3), Inches(4.55), Inches(4.15), AMBER_WASH)
    add_text(slide, Inches(8.45), Inches(2.58), Inches(4.0), Inches(0.9), "A lista vai inteira", SERIF, 24, INK)
    add_text(
        slide,
        Inches(8.45),
        Inches(3.6),
        Inches(3.95),
        Inches(2.3),
        "O primeiro nome pode já estar em outra corrida. O segundo e o terceiro já vêm na mesma mensagem.",
        SANS,
        16,
        BODY,
    )

    notes(
        slide,
        "Política NearestETA: ordena por ETA, depois por distância, depois por id. "
        "exclude_driver_ids tira da lista quem já sumiu nesta corrida. "
        "Sem candidato: match-failed no_candidate. Ninguém alcançável: none_reachable. "
        "O ack no Kafka só sai depois de publicar o resultado. Redis deduplica event_id por 1 hora.",
    )


def slide_routing(prs):
    slide = blank(prs)
    light_chrome(slide, 10)
    kicker(slide, "APLICAÇÃO  04")
    title(slide, "routing-service")
    add_text(
        slide,
        Inches(0.64),
        Inches(1.58),
        Inches(12),
        Inches(0.4),
        "Recebe coordenadas. Devolve tempo e distância. Não conhece corrida nem motorista.",
        SANS,
        16,
        SLATE,
    )

    points = [
        ("Mapa na memória", "Cada réplica lê o mesmo arquivo na subida. Depois disso o mapa não muda."),
        ("Cresce com a CPU", "Começa em duas e vai até quatro. Não há estado para copiar entre elas."),
        ("Se lotar, tenta de novo", "Há um limite de buscas ao mesmo tempo. A réplica cheia pede para tentar outra."),
    ]
    x = 0.62
    for name, body in points:
        card(slide, Inches(x), Inches(2.25), Inches(3.95), Inches(4.35), WHITE)
        rect(slide, Inches(x + 0.32), Inches(2.58), Inches(0.46), Inches(0.06), CORAL)
        add_text(slide, Inches(x + 0.32), Inches(2.9), Inches(3.3), Inches(1.15), name, SERIF, 24, INK)
        add_text(slide, Inches(x + 0.32), Inches(4.25), Inches(3.3), Inches(1.9), body, SANS, 15, BODY)
        x += 4.15

    notes(
        slide,
        "Boot monta o grafo CSR, o grafo reverso e a KD-Tree dos nós para snapping. readyz só fica verde depois disso. "
        "Pool de buffers esgotado responde 503 com Retry-After. Não usa Redis nem banco. Não há Kafka. "
        "O único chamador é o matching, pela rede interna.",
    )


def slide_dijkstra(prs):
    slide = blank(prs)
    light_chrome(slide, 11)
    kicker(slide, "ALGORITMO")
    title(slide, "Uma busca para todos", size=32)
    add_text(
        slide,
        Inches(0.64),
        Inches(1.62),
        Inches(7.1),
        Inches(0.4),
        "Dijkstra. Uma busca, não uma por motorista.",
        SANS,
        15,
        SLATE,
    )

    ideas = [
        ("01", "Começa no embarque", "O cálculo sai do ponto de embarque e caminha pelas ruas até alcançar cada motorista."),
        ("02", "O custo é o tempo", "Uma rua mais curta pode ser mais lenta. Ganha quem chega antes, não quem está menos metros longe."),
        ("03", "Para cedo", "Assim que todos os candidatos foram alcançados. Quem ficou de fora volta como sem caminho."),
    ]
    y = 2.15
    for number, name, body in ideas:
        add_text(slide, Inches(0.62), Inches(y), Inches(1.0), Inches(0.45), number, SERIF, 22, AMBER)
        add_text(slide, Inches(1.7), Inches(y), Inches(5.5), Inches(0.4), name, SERIF, 22, INK)
        add_text(slide, Inches(1.7), Inches(y + 0.48), Inches(5.6), Inches(0.8), body, SANS, 15, BODY)
        y += 1.45

    # Grafo pequeno: embarque à direita, motoristas à esquerda, um solto.
    ox, oy = 8.15, 2.35
    nodes = {
        "d1": (0.2, 0.35, TEAL),
        "d2": (0.15, 1.85, TEAL),
        "d3": (0.35, 3.55, CORAL),
        "i1": (1.55, 0.85, RGBColor(0xC5, 0xD0, 0xC8)),
        "i2": (1.7, 2.35, RGBColor(0xC5, 0xD0, 0xC8)),
        "pin": (3.15, 1.55, AMBER),
    }
    edges = [("d1", "i1"), ("d2", "i1"), ("d2", "i2"), ("i1", "pin"), ("i2", "pin")]
    centers = {}
    for name, (x, y, _) in nodes.items():
        centers[name] = (ox + x + 0.16, oy + y + 0.16)
    for u, v in edges:
        x1, y1 = centers[u]
        x2, y2 = centers[v]
        line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        line.line.color.rgb = RGBColor(0xC8, 0xBE, 0xAE)
        line.line.width = Pt(1.75)
    for name, (x, y, color) in nodes.items():
        d = 0.38 if name == "pin" else 0.32
        oval(slide, Inches(ox + x), Inches(oy + y), Inches(d), color)

    add_text(slide, Inches(8.2), Inches(2.05), Inches(2.2), Inches(0.24), "motoristas", SANS, 12, SLATE)
    add_text(slide, Inches(10.75), Inches(4.42), Inches(1.5), Inches(0.24), "embarque", SANS, 12, AMBER_DEEP, align=PP_ALIGN.CENTER)
    add_text(slide, Inches(8.2), Inches(6.38), Inches(4.6), Inches(0.4), "O ponto coral não tem rua até o embarque.", SANS, 13, SLATE)

    notes(
        slide,
        "É um Dijkstra no grafo reverso, peso duration_s, distância acumulada junto. "
        "Uma busca a partir do nó do embarque, não K rotas. Para quando as origens foram assentadas ou a distância passa de max_radius_m. "
        "Coordenada vira o nó mais próximo por snapping. O nome na resposta é dijkstra-one-to-many.",
    )


def slide_fallback(prs):
    slide = blank(prs)
    light_chrome(slide, 12)
    kicker(slide, "PLANO B")
    title(slide, "Se o mapa não responder", size=32)
    add_text(
        slide,
        Inches(0.64),
        Inches(1.62),
        Inches(7),
        Inches(0.7),
        "A corrida continua. O tempo passa a ser uma estimativa.",
        SANS,
        16,
        SLATE,
    )

    points = [
        ("Linha reta, não rua", "A distância direta entre o motorista e o embarque, dividida por 30 km/h."),
        ("Quando entra", "O routing caiu, ou falhou várias vezes seguidas. Aí o matching nem insiste."),
        ("Fica visível", "A resposta marca que aquele tempo não veio do mapa."),
    ]
    y = 2.5
    for name, body in points:
        add_text(slide, Inches(0.64), Inches(y), Inches(6.8), Inches(0.34), name, SANS, 16, INK, bold=True)
        add_text(slide, Inches(0.64), Inches(y + 0.38), Inches(6.8), Inches(0.7), body, SANS, 15, BODY)
        y += 1.25

    card(slide, Inches(8.15), Inches(2.15), Inches(4.55), Inches(4.3), WHITE)
    add_text(slide, Inches(8.45), Inches(2.5), Inches(4.0), Inches(0.35), "A CONTA", SANS, 12, AMBER_DEEP, bold=True, tracking=160)
    add_lines(
        slide,
        Inches(8.45),
        Inches(3.15),
        Inches(3.95),
        Inches(2.2),
        ["linha reta", "÷  30 km/h"],
        SERIF,
        32,
        INK,
        spacing=1.15,
    )
    add_text(slide, Inches(8.45), Inches(5.55), Inches(3.95), Inches(0.5), "Nome do cálculo: Haversine.", SANS, 14, SLATE)

    notes(
        slide,
        "HAVERSINE_SPEED_MPS padrão é 8,33 m/s, 30 km/h. engine haversine-fallback. "
        "Falha de routing não vira NO_DRIVER_AVAILABLE. O circuit breaker abre após 5 falhas seguidas e entra em meia-abertura em 10 segundos.",
    )


def slide_infra(prs):
    slide = blank(prs)
    light_chrome(slide, 13)
    kicker(slide, "INFRAESTRUTURA")
    title(slide, "O que fica em volta", size=32)

    items = [
        (AMBER, "Kubernetes", "A escala", "Ride e driver ficam em um pod. Matching e routing crescem quando a CPU passa de 60%."),
        (TEAL, "Kafka", "O aviso", "A busca começa depois do aceite. Um serviço não fica preso esperando o outro."),
        (INK_SOFT, "PostgreSQL", "O registro", "Corrida e cadastro do motorista. É o que não pode sumir."),
        (CORAL, "Redis", "O atalho", "Reserva rápida e pedido repetido. Expira sozinho. Não é a fonte da verdade."),
    ]
    x = 0.62
    for color, name, role, body in items:
        card(slide, Inches(x), Inches(2.15), Inches(2.95), Inches(4.45), WHITE)
        oval(slide, Inches(x + 0.26), Inches(2.52), Inches(0.16), color)
        add_text(slide, Inches(x + 0.54), Inches(2.42), Inches(2.2), Inches(0.42), name, SERIF, 22, INK)
        add_text(slide, Inches(x + 0.26), Inches(3.05), Inches(2.45), Inches(0.36), role, SANS, 14, color, bold=True)
        add_text(slide, Inches(x + 0.26), Inches(3.6), Inches(2.45), Inches(2.5), body, SANS, 14, BODY)
        x += 3.15

    notes(
        slide,
        "No cluster: ride 1, driver 1, matching de 1 a 3, routing de 2 a 4. "
        "O autoscaler olha 60% de CPU e sobe um pod a cada 30 segundos. "
        "Kafka: rides por 7 dias, localização por 5 minutos, presença por 1 hora. "
        "Redis sem persistência: só TTL. Postgres 17 com schemas ride e driver separados.",
    )


def slide_where(prs):
    slide = blank(prs)
    light_chrome(slide, 14)
    kicker(slide, "ONDE CADA DADO MORA")
    title(slide, "O que sobrevive a uma queda", size=32)

    rows = [
        (AMBER, "Corrida", "PostgreSQL", "Se o ride cair, o pedido continua no banco. O aviso pendente é publicado de novo."),
        (TEAL, "Posição agora", "Memória + Kafka", "Pode sumir por segundos. Os últimos 5 minutos de GPS reconstroem o índice."),
        (CORAL, "Mapa de ruas", "Arquivo, em cada réplica", "Uma réplica fora não para o sistema. A outra tem o mesmo mapa."),
        (INK_SOFT, "Atalhos", "Redis, com prazo", "Se o Redis cair, a corrida e o cadastro seguem. O banco segura a verdade."),
    ]
    y = 1.9
    for color, name, where, body in rows:
        card(slide, Inches(0.62), Inches(y), Inches(12.1), Inches(1.15), WHITE)
        oval(slide, Inches(0.9), Inches(y + 0.48), Inches(0.18), color)
        add_text(slide, Inches(1.28), Inches(y + 0.16), Inches(3.15), Inches(0.36), name, SANS, 16, INK, bold=True)
        add_text(slide, Inches(1.28), Inches(y + 0.56), Inches(3.15), Inches(0.36), where, SANS, 13, color, bold=True)
        add_text(slide, Inches(4.6), Inches(y + 0.22), Inches(7.8), Inches(0.72), body, SANS, 15, BODY, anchor=MSO_ANCHOR.MIDDLE)
        y += 1.28

    notes(
        slide,
        "Ride e driver: uma instância. O índice espacial do driver vive naquele processo. "
        "Matching de 1 a 3, uma cópia por partição do pedido. Routing de 2 a 4, sem estado. "
        "Redis guarda idempotência por 24 h, reserva por 30 s, última posição por 60 s e dedup de evento por 1 h. "
        "Nada disso é a garantia final.",
    )


def slide_failure(prs):
    slide = blank(prs)
    light_chrome(slide, 15)
    kicker(slide, "QUANDO ALGO REPETE OU FALHA")
    title(slide, "O sistema não faz duas vezes", size=32)

    items = [
        ("Mesmo pedido", "A mesma chave devolve a mesma corrida. Não nasce um segundo pedido."),
        ("Mesma mensagem", "Chegar de novo não atribui outra vez, nem cria outro motorista."),
        ("Motorista some", "Ele sai da corrida. A busca recomeça sem ele, até o limite de tentativas."),
        ("Mapa indisponível", "O tempo cai para a linha reta. A atribuição continua."),
    ]
    positions = [(0.62, 1.95), (6.95, 1.95), (0.62, 4.25), (6.95, 4.25)]
    for (name, body), (x, y) in zip(items, positions):
        card(slide, Inches(x), Inches(y), Inches(5.75), Inches(2.1), WHITE)
        rect(slide, Inches(x + 0.32), Inches(y + 0.32), Inches(0.42), Inches(0.06), AMBER)
        add_text(slide, Inches(x + 0.32), Inches(y + 0.52), Inches(5.1), Inches(0.45), name, SERIF, 22, INK)
        add_text(slide, Inches(x + 0.32), Inches(y + 1.1), Inches(5.15), Inches(0.75), body, SANS, 15, BODY)

    notes(
        slide,
        "Idempotency-Key no POST /rides. attempt descarta ranking de tentativa antiga. seq descarta GPS atrasado. "
        "rides.released só libera se o ride_id for o da reserva atual. "
        "Segundo lost da mesma corrida não transiciona de novo, pela version otimista.",
    )


def slide_close(prs):
    slide = blank(prs)
    dark_chrome(slide, 16)
    rect(slide, 0, 0, Inches(0.08), H, AMBER)
    add_text(slide, Inches(0.78), Inches(1.7), Inches(8), Inches(0.3), "NEARBY-DRIVERS", SANS, 13, AMBER, bold=True, tracking=220)
    add_lines(
        slide,
        Inches(0.74),
        Inches(2.2),
        Inches(11),
        Inches(3.4),
        ["Quem está livre.", "Quem chega antes.", "Em quanto tempo."],
        SERIF,
        40,
        WHITE,
        spacing=1.15,
    )

    notes(
        slide,
        "Fecha repetindo a promessa do começo. Se perguntarem o que o passageiro vê: o motorista atribuído, o ETA e a distância. "
        "Se o mapa estava fora, o ETA vem marcado como estimativa em linha reta.",
    )


def build(path):
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    prs.core_properties.title = "nearby-drivers"
    prs.core_properties.subject = "Como o sistema encontra o motorista mais próximo"
    prs.core_properties.category = "Apresentação"

    slide_cover(prs)
    slide_promise(prs)
    slide_apps(prs)
    slide_flow(prs)
    slide_ride(prs)
    slide_states(prs)
    slide_driver(prs)
    slide_kdtree(prs)
    slide_matching(prs)
    slide_routing(prs)
    slide_dijkstra(prs)
    slide_fallback(prs)
    slide_infra(prs)
    slide_where(prs)
    slide_failure(prs)
    slide_close(prs)

    prs.save(path)


if __name__ == "__main__":
    out = "slides/nearby-drivers.pptx"
    build(out)
    print(out)
